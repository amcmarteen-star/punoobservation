"""
Excel export of reforestation sites, in the DENR contract profile layout.

The year sheets copy the DENR NGP contract profile workbook, the same file
the Reference Dataset page imports:

    row 1   REGION 1
    row 2   CONTRACT PROFILE
    row 3   the DENR column headers
    row 4+  one site per row, one sheet per contract year

The header text is exactly what the importer reads, so an exported file can
be imported again.

Extra sheets hold what only this system records: monitoring reports and
requests. The importer skips any sheet whose row 3 lacks the DENR site
columns, so those sheets never interfere with a re-import.

Nothing here decides who may see what. The route passes in sites, reports
and requests it has already scoped to the user.
"""

import io
from collections import defaultdict
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.extensions import db
from app.models import ReforestationRecord, TreeSpecie
from app.services.monitoring import DENR_THRESHOLD
from app.utils.timeutil import manila_now

# Same text as admin.PLACEHOLDER_SPECIES_NAME: the species the importer
# gives a planting record before real species are matched. It is not a
# species anyone planted, so it is never exported as one.
PLACEHOLDER_SPECIES_NAME = "Unspecified (Pending Match)"

REGION_TITLE = "REGION 1"
SHEET_TITLE = "CONTRACT PROFILE"
NO_YEAR_SHEET = "No Year Contracted"
HEADER_ROW = 3

DATE_FMT = "mmmm d, yyyy"
DATETIME_FMT = "mmmm d, yyyy h:mm AM/PM"
PERCENT_FMT = "0.00%"
COUNT_FMT = "#,##0"
MONEY_FMT = "#,##0.00"

# The narrowest column that shows a formatted date without Excel falling
# back to "###". The longest values are "September 30, 2026" and
# "September 30, 2026 12:00 PM".
DATE_WIDTH = 20
DATETIME_WIDTH = 28

# Styling taken from the DENR workbook: bold size 12 titles and headers,
# a light green header fill, centred wrapped header text, thin borders.
TITLE_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True, size=12)
HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
NOTE_FONT = Font(italic=True, size=10, color="7F6000")
_THIN = Side(style="thin")
HEADER_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# (header, column width, number format)
# Header text is exactly what the importer reads. Order follows the DENR
# master sheet, limited to the columns this system stores.
CONTRACT_COLUMNS = [
    ("PENRO", 13, None),
    ("IMPLEMENTING CENRO", 14.11, None),
    ("CONGRESSIONAL DISTRICT", 18.79, None),
    ("BARANGAY", 13, None),
    ("MUNICIPALITY", 13, None),
    ("IDENTIFY IF WATERSHED, PA OR REGULAR", 13, None),
    ("MAJOR LAND-USE LOCATION (Name of Watershed,PA or NA)", 26.77, None),
    ("ZONE (PRODUCTION OR PROTECTION)", 13, None),
    ("SITE CODE", 26, None),
    ("NAME OF PARTNER/CONTRACTOR (GROUP OR INDIVIDUAL)", 38, None),
    ("NAME OF GROUP CONTACT PERSON OF THE PARTNER GROUP", 21, None),
    ("LOA/CONTRACT CODE (YEAR-CENRO CODE-3 DIGITS)", 25.67, None),
    ("YEAR CONTRACTED", 13, None),
    ("AREA CONTRACTED (HA)", 13, None),
    ("ACTUAL AREA PLANTED (HA)", 13, None),
    ("NO. OF SEEDLINGS TO BE PLANTED", 13, COUNT_FMT),
    ("NO. OF SEEDLINGS PLANTED", 13, COUNT_FMT),
    ("COMPONENT", 13, None),
    ("COMMODITY (if mixed,specify all commodities within the contracted Site)", 23.88, None),
    ("SPECIES TO BE PLANTED", 30, None),
    ("PROJECT COST (3 YEARS)", 14, MONEY_FMT),
    ("DATE OF EXECUTION OF CONTRACT (MM/DD/YR) 15 DAYS UPON ISSUANCE OF NTP)", 18, DATE_FMT),
    ("DATE OF EXPIRY OF CONTRACT (END DATE AT YEAR 3)", 22.76, DATE_FMT),
    ("RETENTION FEE at 3rd YEAR (Amount paid)", 14, MONEY_FMT),
    ("RETENTION FEE at 3rd YEAR (Date paid)", 18, DATE_FMT),
    ("AMOUNT STILL IN CIP", 14, MONEY_FMT),
    ("AMOUNT UNDER LAND IMPROVEMENT", 14, MONEY_FMT),
    ("DATE OF PERFORMANCE VALIDATION REPORT (IAC REPORT)", 19.67, DATE_FMT),
    ("SURVIVAL RATE ON THE 3RD YEAR", 13, PERCENT_FMT),
]

# Headers on the extra sheets are title case on purpose. The importer
# matches the upper-case DENR names, so it never mistakes these sheets
# for site sheets.
REPORT_COLUMNS = [
    ("Report ID", 10, None),
    ("Site Code", 22, None),
    ("Site Name", 34, None),
    ("Barangay", 16, None),
    ("Municipality", 16, None),
    ("Implementing CENRO", 16, None),
    ("Monitoring Date", 18, DATE_FMT),
    ("Submitted", 28, DATE_FMT),
    ("Field Officer", 16, None),
    ("Plots Recorded", 10, None),
    ("Plot Size (sq m)", 10, None),
    ("Seedlings Counted", 12, COUNT_FMT),
    ("Estimated Survivors", 12, COUNT_FMT),
    ("Survival Rate", 11, PERCENT_FMT),
    (f"Meets {DENR_THRESHOLD:g}%", 10, None),
    ("Captured Boundary (ha)", 12, "0.00"),
    ("Boundary Publication", 18, None),
    ("Reviewed By", 16, None),
    ("Date Reviewed", 18, DATE_FMT),
]

REQUEST_COLUMNS = [
    ("Request ID", 10, None),
    ("Date Submitted", 18, DATE_FMT),
    ("Requester", 16, None),
    ("Barangay", 16, None),
    ("Municipality", 16, None),
    ("Request Type", 20, None),
    ("Status", 14, None),
    ("Proposed Area (ha)", 12, "0.00"),
    ("Drawn Area (ha)", 12, "0.00"),
    ("Land Ownership", 22, None),
    ("Land Cover", 22, None),
    ("Planting Density (per ha)", 12, COUNT_FMT),
    ("Estimated Seedlings", 12, COUNT_FMT),
    ("Latitude", 12, "0.000000"),
    ("Longitude", 12, "0.000000"),
    ("Reviewed By", 16, None),
    ("Date Reviewed", 18, DATE_FMT),
    ("Site Created", 34, None),
    ("Site Code", 22, None),
]


# ----------------------------------------------------------------------
# public
# ----------------------------------------------------------------------

def build_sites_workbook(sites, reports, requests, exported_by, scope):
    """
    The whole export as .xlsx bytes in a BytesIO, ready for send_file.

    sites, reports and requests must already be limited to what the user
    may see; this function exports everything it is given.
    """
    site_ids = [s.site_id for s in sites]
    records = _records_by_site(site_ids)
    species = _species_by_site(site_ids)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    by_year = defaultdict(list)
    for site in sites:
        by_year[site.year_contracted].append(site)

    for year in sorted(y for y in by_year if y is not None):
        _write_contract_sheet(wb.create_sheet(str(year)), by_year[year],
                              records, species)

    if None in by_year:
        _write_contract_sheet(
            wb.create_sheet(NO_YEAR_SHEET), by_year[None], records, species,
            note=("Sites with no contract year yet, such as sites created "
                  "from requests. A site with no site code is added as a "
                  "new site each time this file is imported."),
        )

    _write_summary(summary, sites, exported_by, scope)
    _write_table(wb.create_sheet("Monitoring Reports"),
                 "APPROVED MONITORING REPORTS", REPORT_COLUMNS,
                 [_report_row(r) for r in reports])
    _write_table(wb.create_sheet("Requests"),
                 "REFORESTATION REQUESTS", REQUEST_COLUMNS,
                 [_request_row(q) for q in requests])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# ----------------------------------------------------------------------
# data
# ----------------------------------------------------------------------

def _records_by_site(site_ids):
    out = defaultdict(list)
    if not site_ids:
        return out
    for rec in ReforestationRecord.query.filter(
            ReforestationRecord.site_id.in_(site_ids)).all():
        out[rec.site_id].append(rec)
    return out


def _species_by_site(site_ids):
    """Real species names per site, comma separated. Placeholder left out."""
    if not site_ids:
        return {}
    rows = (
        db.session.query(ReforestationRecord.site_id, TreeSpecie.specie_name)
        .join(TreeSpecie, TreeSpecie.tree_id == ReforestationRecord.tree_id)
        .filter(ReforestationRecord.site_id.in_(site_ids),
                TreeSpecie.specie_name != PLACEHOLDER_SPECIES_NAME)
        .distinct()
        .all()
    )
    names = defaultdict(set)
    for site_id, name in rows:
        names[site_id].add(name)
    return {sid: ", ".join(sorted(n)) for sid, n in names.items()}


def _planting_figures(records):
    """
    (target, planted, validation date, survival fraction) for one site.

    Seedling figures are summed over the site's planting records. Survival
    comes from the most recently validated record that has one, stored as
    a percentage (86.0) and written as the fraction Excel keeps behind a
    percentage cell (0.86), which is also what the importer expects.
    """
    targets = [r.target_quantity for r in records if r.target_quantity is not None]
    planted = [r.actual_quantity_planted for r in records
               if r.actual_quantity_planted is not None]
    surveyed = [r for r in records if r.survival_rate is not None]
    latest = max(surveyed,
                 key=lambda r: r.date_validated or r.date_planted or date.min,
                 default=None)
    return (
        sum(targets) if targets else None,
        sum(planted) if planted else None,
        latest.date_validated if latest else None,
        latest.survival_rate / 100.0 if latest else None,
    )


def _contract_row(site, records, species):
    target, planted, validated, survival = _planting_figures(records)
    loc = site.location
    return [
        site.penro,
        site.cenro,
        site.congressional_district,
        loc.barangay if loc else None,
        loc.municipality if loc else None,
        site.land_classification,
        site.major_land_use,
        site.zone_type,
        site.site_code,
        site.organization.organization_name if site.organization else None,
        site.contact_person,
        site.loa_contract_code,
        site.year_contracted,
        site.area_contracted_ha,
        site.area_size_ha,
        target,
        planted,
        site.component,
        site.commodity,
        species or None,
        site.project_cost_3yr,
        site.date_contract_executed,
        site.date_contract_expiry,
        site.retention_fee_amount_paid,
        site.retention_fee_date_paid,
        site.amount_still_in_cip,
        site.amount_under_land_improvement,
        validated,
        survival,
    ]


def _report_row(r):
    site = r.site
    loc = site.location if site else None
    rate = r.survival_rate
    return [
        r.report_id,
        site.site_code if site else None,
        site.site_name if site else None,
        loc.barangay if loc else None,
        loc.municipality if loc else None,
        site.cenro if site else None,
        r.monitoring_date,
        r.submitted_at,
        r.officer.username if r.officer else None,
        r.plots_recorded,
        r.plot_size_sqm,
        r.total_counted,
        r.estimated_survivors,
        rate / 100.0 if rate is not None else None,
        ("Yes" if rate >= DENR_THRESHOLD else "No") if rate is not None else None,
        r.captured_area_ha,
        r.publication_status,
        r.reviewer.username if r.reviewer else None,
        r.date_reviewed,
    ]


def _request_row(q):
    loc = q.location
    site = q.created_site
    return [
        q.request_id,
        q.date_submitted,
        q.requester.username if q.requester else None,
        loc.barangay if loc else None,
        loc.municipality if loc else None,
        q.request_type,
        q.status,
        q.proposed_area_ha,
        q.boundary_area_ha,
        q.land_ownership,
        q.land_cover,
        q.planting_density_per_ha,
        q.estimated_seedlings,
        q.latitude,
        q.longitude,
        q.reviewer.username if q.reviewer else None,
        q.date_reviewed,
        site.site_name if site else None,
        site.site_code if site else None,
    ]


# ----------------------------------------------------------------------
# sheets
# ----------------------------------------------------------------------

def _excel_value(value):
    """Excel has no time zones, so aware datetimes lose theirs."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _write_header(ws, columns, row=HEADER_ROW):
    for col, (header, width, _fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = HEADER_BORDER
        ws.column_dimensions[get_column_letter(col)].width = width


def _widen(ws, col, width):
    """Make a column at least this wide. Never narrows it."""
    letter = get_column_letter(col)
    if (ws.column_dimensions[letter].width or 0) < width:
        ws.column_dimensions[letter].width = width


def _write_rows(ws, columns, rows, first_row=HEADER_ROW + 1):
    for r, values in enumerate(rows, start=first_row):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=_excel_value(value))
            fmt = columns[col - 1][2]
            if fmt and value is not None:
                # A datetime also shows its time, so it needs the wider
                # format and a wider column. Too narrow shows "###".
                if fmt == DATE_FMT and isinstance(value, datetime):
                    fmt = DATETIME_FMT
                    _widen(ws, col, DATETIME_WIDTH)
                elif fmt == DATE_FMT:
                    _widen(ws, col, DATE_WIDTH)
                cell.number_format = fmt


def _write_contract_sheet(ws, sites, records, species, note=None):
    ws["A1"] = REGION_TITLE
    ws["A1"].font = TITLE_FONT
    ws["A2"] = SHEET_TITLE
    ws["A2"].font = TITLE_FONT
    if note:
        ws["D1"] = note
        ws["D1"].font = NOTE_FONT

    _write_header(ws, CONTRACT_COLUMNS)
    ws.row_dimensions[HEADER_ROW].height = 95

    _write_rows(ws, CONTRACT_COLUMNS, [
        _contract_row(s, records.get(s.site_id, []), species.get(s.site_id))
        for s in sites
    ])

    # Same as the DENR 2021 sheet: titles and headers stay put, and so do
    # the columns up to SITE CODE while scrolling right.
    ws.freeze_panes = "J4"


def _write_table(ws, title, columns, rows):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    _write_header(ws, columns)
    ws.row_dimensions[HEADER_ROW].height = 45
    _write_rows(ws, columns, rows)
    if not rows:
        ws.cell(row=HEADER_ROW + 1, column=1, value="None in this export.").font = NOTE_FONT
    ws.freeze_panes = "A4"


def _write_summary(ws, sites, exported_by, scope):
    """
    The DENR summary pivot: actual area planted, per CENRO and contract
    year, with Total Result on both axes. Then who exported what, and when.
    """
    bold = Font(bold=True)
    no_year = "No year"

    years = sorted({s.year_contracted for s in sites if s.year_contracted is not None})
    columns = years + ([no_year] if any(s.year_contracted is None for s in sites) else [])
    cenros = sorted({s.cenro or "Unassigned" for s in sites})

    area = defaultdict(float)
    for s in sites:
        key = s.year_contracted if s.year_contracted is not None else no_year
        area[(s.cenro or "Unassigned", key)] += float(s.area_size_ha or 0)

    ws["A1"] = "SUM of ACTUAL AREA PLANTED (HA)"
    ws["B1"] = "YEAR CONTRACTED"
    ws["A2"] = "IMPLEMENTING CENRO"
    for i, key in enumerate(columns, start=2):
        ws.cell(row=2, column=i, value=key)
    total_col = len(columns) + 2
    ws.cell(row=2, column=total_col, value="Total Result")
    for cell in (*ws[1], *ws[2]):
        cell.font = bold

    row = 3
    for cenro in cenros:
        ws.cell(row=row, column=1, value=cenro)
        for i, key in enumerate(columns, start=2):
            value = area.get((cenro, key))
            ws.cell(row=row, column=i, value=round(value, 3) if value else None)
        ws.cell(row=row, column=total_col,
                value=round(sum(area.get((cenro, k), 0) for k in columns), 3))
        row += 1

    ws.cell(row=row, column=1, value="Total Result").font = bold
    for i, key in enumerate(columns, start=2):
        ws.cell(row=row, column=i,
                value=round(sum(area.get((c, key), 0) for c in cenros), 3)).font = bold
    ws.cell(row=row, column=total_col,
            value=round(sum(area.values()), 3)).font = bold

    ws.column_dimensions["A"].width = 26
    for i in range(2, total_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13

    exported_on = manila_now()
    info = [
        ("Exported by", exported_by),
        ("Exported on", exported_on),
        ("Scope", scope),
        ("Sites", len(sites)),
        ("About this file",
         "Each year sheet follows the DENR contract profile layout and can "
         "be imported again on the Reference Dataset page. Sites are matched "
         "by SITE CODE, so a site without one is added as a new site on "
         "every import."),
    ]
    row += 2
    for label, value in info:
        ws.cell(row=row, column=1, value=label).font = bold
        cell = ws.cell(row=row, column=2, value=value)
        if isinstance(value, datetime):
            cell.number_format = DATETIME_FMT
            _widen(ws, 2, DATETIME_WIDTH)
        row += 1
